"""Charger une soirée depuis le classeur du soir.

Le point délicat n'est pas la lecture du fichier — elle est éprouvée côté unité — mais ce que
l'import a le droit de récrire. La règle tient en une phrase : **il ne récrit rien de ce qui a
été pointé**. C'est la leçon de l'import d'inventaire, où un fichier qui reprend tout annulait à
chaque dépôt le travail fait à l'écran.

On vérifie donc, dans l'ordre : qu'un dépôt remplit une soirée absente, qu'un second dépôt du même
fichier ne change plus rien, qu'une étape déjà pointée survit au dépôt, et qu'un intitulé inconnu
est montré plutôt que perdu.
"""

from datetime import datetime
from io import BytesIO
from typing import Any

import openpyxl
from httpx import AsyncClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from tests.integration.conftest import creer_utilisateur, entetes

_JOURNEE = "2026-05-12"


def _classeur(lignes: list[list[Any]]) -> bytes:
    """Un classeur à la forme de la banque : intitulé en D, Start/End/Notes en E, F et G."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    assert feuille is not None
    for colonne, valeur in enumerate(["AFG-BANK-MALI", "Start", "End", "Notes"], start=4):
        feuille.cell(2, colonne, valeur)
    for decalage, ligne in enumerate(lignes):
        for colonne, valeur in enumerate(ligne, start=4):
            if valeur is not None:
                feuille.cell(3 + decalage, colonne, valeur)
    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


#: Un extrait fidèle du rapport réel : ponctuation flottante comprise (« CSSJOBBR--- »).
_NUIT = [
    ["INTEGRATION FICHIER CARTHAGO", "17H55", "19H12", "COMPLETED"],
    ["Backup before EOM", "X", "X", "X"],
    ["System Date", datetime(2026, 5, 12), None, "Complete"],
    ["Check Batch CSSJOBBR--- AC-DAHOFF", "19H45", "19H45", "Complete"],
    ["Date Check", "07H30", "07H31", "Le controle a tourne au matin."],
]


async def _deposer(client: AsyncClient, uid: str, contenu: bytes) -> dict[str, Any]:
    r = await client.post(
        "/import/eod",
        headers=entetes(uid),
        files={"fichier": ("eod.xlsx", contenu, "application/vnd.ms-excel")},
    )
    assert r.status_code == 200, r.text
    return dict(r.json())


async def _etapes(session: AsyncSession, soiree_id: str) -> dict[str, Any]:
    lignes = (
        await session.execute(
            text(
                "SELECT libelle, statut, debut, fin, valeur FROM core.eod_etape "
                "WHERE activite_id = cast(:a as uuid)"
            ),
            {"a": soiree_id},
        )
    ).mappings()
    return {str(ligne["libelle"]): dict(ligne) for ligne in lignes}


async def test_le_depot_ouvre_la_soiree_et_la_remplit(
    client: AsyncClient, session: AsyncSession
) -> None:
    agent = await creer_utilisateur(session, email="prod.eod@afgbank.ml", profil="ADMIN")

    rapport = await _deposer(client, agent, _classeur(_NUIT))
    assert rapport["creee"] is True
    assert rapport["journee"] == _JOURNEE
    assert rapport["etapes_renseignees"] == 5
    assert rapport["observations_ajoutees"] == 1
    assert rapport["lignes_non_reconnues"] == []

    etapes = await _etapes(session, rapport["soiree_id"])
    # Le verdict se déduit des heures ; « X » dit que l'étape ne s'appliquait pas ce soir-là.
    assert etapes["Intégration fichier CARTHAGO"]["statut"] == "Complété"
    assert etapes["Backup before EOM"]["statut"] == "Non applicable"
    # La ligne à valeur relève ce qu'affichait l'écran, pas l'instant du regard.
    assert etapes["System Date"]["valeur"] == "12/05/2026"
    # L'intitulé se reconnaît malgré la ponctuation du classeur (« CSSJOBBR--- »).
    assert etapes["Check Batch CSSJOBBR --- AC-DAHOFF"]["statut"] == "Complété"
    # Avant midi : la soirée a franchi minuit, l'étape appartient au lendemain.
    assert etapes["Date Check"]["debut"].date().isoformat() == "2026-05-13"


async def test_un_second_depot_du_meme_fichier_ne_change_rien(
    client: AsyncClient, session: AsyncSession
) -> None:
    """Sinon le journal doublerait le récit de la nuit à chaque dépôt."""
    agent = await creer_utilisateur(session, email="prod.eod2@afgbank.ml", profil="ADMIN")
    premier = await _deposer(client, agent, _classeur(_NUIT))
    second = await _deposer(client, agent, _classeur(_NUIT))

    assert second["creee"] is False
    assert second["reference"] == premier["reference"]
    assert second["etapes_renseignees"] == 0
    assert second["etapes_conservees"] == 5
    assert second["observations_ajoutees"] == 0


async def test_une_etape_deja_pointee_survit_au_depot(
    client: AsyncClient, session: AsyncSession
) -> None:
    """L'écran fait foi : le classeur complète ce qui manque, il ne corrige pas ce qui est là."""
    agent = await creer_utilisateur(session, email="prod.eod3@afgbank.ml", profil="ADMIN")
    ouverture = await client.post("/eod", headers=entetes(agent), json={"journee": _JOURNEE})
    assert ouverture.status_code == 201, ouverture.text
    soiree = str(ouverture.json()["id"])

    etapes = await _etapes(session, soiree)
    identifiant = (
        await session.execute(
            text(
                "SELECT id::text FROM core.eod_etape WHERE activite_id = cast(:a as uuid) "
                "AND libelle = 'Intégration fichier CARTHAGO'"
            ),
            {"a": soiree},
        )
    ).scalar_one()
    r = await client.patch(
        f"/eod/{soiree}/etapes/{identifiant}",
        headers=entetes(agent),
        json={"statut": "Complété", "debut": "2026-05-12T20:00:00+00:00"},
    )
    assert r.status_code == 200, r.text

    rapport = await _deposer(client, agent, _classeur(_NUIT))
    assert rapport["etapes_conservees"] == 1
    etapes = await _etapes(session, soiree)
    # 20H00 saisi à l'écran, et non 17H55 lu dans le fichier.
    assert etapes["Intégration fichier CARTHAGO"]["debut"].hour == 20


async def _type_de(session: AsyncSession, soiree_id: str) -> str | None:
    return await session.scalar(
        text(
            "SELECT c.libelle FROM core.activite a "
            "JOIN core.categorie c ON c.id = a.categorie_id WHERE a.id = cast(:i as uuid)"
        ),
        {"i": soiree_id},
    )


async def test_le_type_se_deduit_de_la_date(client: AsyncClient, session: AsyncSession) -> None:
    """Un dépôt ne laisse pas la colonne « Type » vide : la date dit déjà ce que la nuit portait."""
    agent = await creer_utilisateur(session, email="prod.eod5@afgbank.ml", profil="ADMIN")

    ordinaire = await _deposer(client, agent, _classeur(_NUIT))
    assert await _type_de(session, ordinaire["soiree_id"]) == "EOD"

    # Le 31 mai clôt le mois : la soirée porte les traitements mensuels en plus.
    fin_de_mois = [["System Date", datetime(2026, 5, 31), None, "Complete"]]
    rapport = await _deposer(client, agent, _classeur(fin_de_mois))
    assert await _type_de(session, rapport["soiree_id"]) == "EOM"


async def test_le_type_se_corrige_a_la_main(client: AsyncClient, session: AsyncSession) -> None:
    """Un arrêté décalé se tranche à la main, sans rouvrir la nuit."""
    agent = await creer_utilisateur(session, email="prod.eod6@afgbank.ml", profil="ADMIN")
    rapport = await _deposer(client, agent, _classeur(_NUIT))
    assert await _type_de(session, rapport["soiree_id"]) == "EOD"

    voulu = await session.scalar(
        text("SELECT id::text FROM core.categorie WHERE module = 'eod' AND code = 'FIN_ANNEE'")
    )
    r = await client.patch(
        f"/eod/{rapport['soiree_id']}/type",
        headers=entetes(agent),
        json={"categorie_id": voulu},
    )
    assert r.status_code == 200, r.text
    assert r.json()["categorie"] == "EOY"
    assert await _type_de(session, rapport["soiree_id"]) == "EOY"


async def test_un_type_etranger_au_module_est_refuse(
    client: AsyncClient, session: AsyncSession
) -> None:
    agent = await creer_utilisateur(session, email="prod.eod7@afgbank.ml", profil="ADMIN")
    rapport = await _deposer(client, agent, _classeur(_NUIT))
    autre = await session.scalar(
        text("SELECT id::text FROM core.categorie WHERE module <> 'eod' LIMIT 1")
    )
    r = await client.patch(
        f"/eod/{rapport['soiree_id']}/type",
        headers=entetes(agent),
        json={"categorie_id": autre},
    )
    assert r.status_code == 400, r.text


async def test_un_intitule_inconnu_est_montre_et_non_perdu(
    client: AsyncClient, session: AsyncSession
) -> None:
    agent = await creer_utilisateur(session, email="prod.eod4@afgbank.ml", profil="ADMIN")
    lignes = [*_NUIT, ["Une etape que le modele ignore", "21H00", "21H05", "Complete"]]

    rapport = await _deposer(client, agent, _classeur(lignes))
    assert rapport["lignes_non_reconnues"] == ["Préparation · Une etape que le modele ignore"]
