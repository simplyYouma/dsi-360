"""Lecture du classeur EOD du soir.

Chaque cas éprouvé ici vient d'une particularité des rapports réels d'AFG Bank Mali (12 et
13/10/2026), et non d'un fichier idéal : c'est le classeur de la banque qu'il faut savoir lire,
avec ses espaces insécables, sa ponctuation flottante et ses « X ».
"""

from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl
import pytest

from dsi360.domain.eod import heure_de_la_nuit
from dsi360.infrastructure.ingestion_eod import ClasseurIllisible, analyser_classeur


def _classeur(lignes: list[list[Any]], entetes: bool = True) -> bytes:
    """Un classeur à la forme de la banque : intitulé en D, Start/End/Notes en E, F et G."""
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    assert feuille is not None
    depart = 2
    if entetes:
        ligne_entete = ["AFG-BANK-MALI", "Start\xa0\xa0", "End\xa0\xa0", "Notes\xa0\xa0"]
        for colonne, valeur in enumerate(ligne_entete, start=4):
            feuille.cell(depart, colonne, valeur)
        depart += 1
    for decalage, ligne in enumerate(lignes):
        for colonne, valeur in enumerate(ligne, start=4):
            if valeur is not None:
                feuille.cell(depart + decalage, colonne, valeur)
    tampon = BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


_NUIT = [
    ["INTEGRATION FICHIER CARTHAGO", "17H55", "19H12", "COMPLETED"],
    ["Backup before EOM", "X", "X", "X"],
    ["System Date", datetime(2026, 10, 13), None, "Complete\xa0\xa0"],
    ["Check Batch CSSJOBBR--- AC-DAHOFF", "19H45", "19H45", "Complete\xa0\xa0"],
    ["PART 1", None, None, None],
    ["Post EOFI_1 for all branch including 000 BAM", "19H46", "20H12", "Complete\xa0\xa0"],
    ["Additionnal Tasks", None, None, None],
    ["Batch Check : EMS_IN", "20H52", "20H52", "The Jobs are started but the date is still 13/10"],
    ["System Date", datetime(2026, 10, 14), None, "Complete\xa0\xa0"],
]


def test_la_journee_comptable_vient_du_premier_system_date() -> None:
    """Et non du nom du fichier : les exemplaires reçus se contredisent déjà là-dessus."""
    classeur = analyser_classeur(_classeur(_NUIT))
    assert classeur.journee == date(2026, 10, 13)


def test_les_sections_se_deduisent_de_la_position() -> None:
    """Les premières lignes ouvrent la soirée ; une ligne sans heure ni note ouvre une section."""
    lignes = analyser_classeur(_classeur(_NUIT)).lignes
    par_libelle = {ligne.libelle: ligne.section for ligne in lignes}
    assert par_libelle["INTEGRATION FICHIER CARTHAGO"] == "Préparation"
    assert par_libelle["Post EOFI_1 for all branch including 000 BAM"] == "PART 1"
    # Seule section dont le nom diffère du déroulé de référence.
    assert par_libelle["Batch Check : EMS_IN"] == "Tâches additionnelles"


def test_une_section_inconnue_traverse_sans_etre_rejetee() -> None:
    """Le vide des colonnes la désigne, pas une liste de noms : « PART 5 » se rangerait seul."""
    lignes = analyser_classeur(
        _classeur([["PART 5", None, None, None], ["Nouvelle etape", "01H00", "01H05", "Complete"]])
    ).lignes
    assert lignes[0].section == "PART 5"


def test_la_croix_marque_une_etape_sans_objet() -> None:
    lignes = {ligne.libelle: ligne for ligne in analyser_classeur(_classeur(_NUIT)).lignes}
    assert lignes["Backup before EOM"].sans_objet is True
    assert lignes["Backup before EOM"].debut is None
    assert lignes["INTEGRATION FICHIER CARTHAGO"].sans_objet is False


def test_un_verdict_n_est_pas_une_observation() -> None:
    """« Complete » dit ce que les heures disent déjà. L'importer noierait les vraies remarques."""
    lignes = {ligne.libelle: ligne for ligne in analyser_classeur(_classeur(_NUIT)).lignes}
    assert lignes["INTEGRATION FICHIER CARTHAGO"].note is None
    assert lignes["Check Batch CSSJOBBR--- AC-DAHOFF"].note is None
    assert lignes["Batch Check : EMS_IN"].note is not None
    assert "13/10" in str(lignes["Batch Check : EMS_IN"].note)


def test_une_ligne_a_valeur_ne_porte_pas_d_heures() -> None:
    """« System Date » relève ce qu'affiche l'écran, pas le moment où on l'a regardé."""
    lignes = [ligne for ligne in analyser_classeur(_classeur(_NUIT)).lignes if ligne.valeur]
    assert [ligne.valeur for ligne in lignes] == [date(2026, 10, 13), date(2026, 10, 14)]
    assert all(ligne.debut is None and ligne.fin is None for ligne in lignes)


def test_un_fichier_sans_les_entetes_est_refuse() -> None:
    """L'écran doit dire pourquoi, plutôt que d'annoncer « 0 ligne importée »."""
    with pytest.raises(ClasseurIllisible):
        analyser_classeur(_classeur([["Une ligne", "01H00", "01H05", "Complete"]], entetes=False))


def test_une_heure_d_avant_midi_appartient_au_lendemain() -> None:
    """La soirée commence l'après-midi du jour comptable et finit au matin du lendemain."""
    journee = date(2026, 10, 13)
    soir = heure_de_la_nuit("20H29", journee)
    matin = heure_de_la_nuit("07H50", journee)
    assert soir is not None and soir.date() == date(2026, 10, 13)
    assert matin is not None and matin.date() == date(2026, 10, 14)


def test_un_recul_dans_les_heures_ne_fait_pas_changer_de_jour() -> None:
    """Le 13/10, « Check Pending Transactions » démarre à 16H00 après une ligne de 17H55 : ce
    contrôle court depuis l'après-midi. Lire ce recul comme un passage de minuit aurait renvoyé la
    ligne au lendemain, et faussé la durée de la nuit."""
    journee = date(2026, 10, 13)
    tardif = heure_de_la_nuit("17H55", journee)
    anterieur = heure_de_la_nuit("16H00", journee)
    assert tardif is not None and anterieur is not None
    assert anterieur.date() == tardif.date() == journee


def test_une_heure_illisible_ne_se_devine_pas() -> None:
    assert heure_de_la_nuit("tout a l'heure", date(2026, 10, 13)) is None
    assert heure_de_la_nuit("25H00", date(2026, 10, 13)) is None
