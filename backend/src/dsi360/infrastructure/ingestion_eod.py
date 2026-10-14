"""Lecture du classeur EOD du soir — celui que la Production remplit depuis toujours.

Le déroulé se pointait dans un tableur avant d'exister ici, et il s'y pointe encore les nuits où
la plateforme n'est pas sous la main. Rien n'oblige alors à retaper vingt-huit lignes à l'écran :
on charge le fichier, et le pointage s'en déduit.

La forme du classeur n'est pas négociable : c'est celle de la banque, relevée sur les rapports du
12 et du 13/10/2026. Une feuille, quatre colonnes utiles — l'intitulé, « Start », « End »,
« Notes » — et des **lignes de section** (« PART 1 », « Additionnal Tasks ») qui ne portent ni
heure ni note. Ce module ne fait que lire et normaliser : il ne décide rien, n'écrit rien, et
rend ce qu'il a compris pour que la couche application le confronte au déroulé de référence.

Trois pièges du fichier réel, qu'on prend en compte ici :

* les intitulés sont ponctués librement — « CSSJOBBR--- AC-DAHOFF » ici, « CSSJOBBR --- AC-DAHOFF »
  dans le modèle — et truffés d'espaces insécables. Ils ne se comparent qu'après normalisation ;
* « X » dit qu'une étape ne s'appliquait pas ce soir-là. Elle est posée indifféremment dans
  « Start », « End » ou « Notes » ;
* la colonne « Notes » sert à deux choses : poser le verdict (« Complete ») ou consigner ce qui
  s'est passé. Seul le second cas est une observation — importer « Complete » comme un commentaire
  remplirait le journal de bruit, et noierait les vraies remarques.
"""

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl

from dsi360.domain.eod import ADDITIONNELLES, PREPARATION
from dsi360.infrastructure.classeur import normaliser

#: Ce que la colonne « Notes » écrit quand elle ne dit rien d'autre que « c'est fait ». Le verdict
#: se déduit déjà des heures : le recopier en observation n'apprendrait rien.
_VERDICTS_MUETS = frozenset({"complete", "completed", "complet", "ok", "done"})

#: La marque d'une étape sans objet ce soir-là, telle qu'elle se pose dans le classeur.
_SANS_OBJET = frozenset({"x", "n/a", "na", "sans objet"})

#: Les en-têtes qui identifient la feuille du rapport. Sans elles, ce n'est pas un rapport EOD.
_COLONNES = ("start", "end", "notes")

#: La seule section dont le nom diffère entre le classeur et le déroulé de référence. Les autres
#: (« PART 1 » à « PART 4 ») sont identiques et traversent sans traduction.
_SECTIONS_TRADUITES = {
    "additionnal tasks": ADDITIONNELLES,  # l'orthographe du classeur de la banque
    "additional tasks": ADDITIONNELLES,
    "taches additionnelles": ADDITIONNELLES,
}

#: Les premières lignes du rapport n'ont pas d'intitulé de section : elles ouvrent la soirée.
_SECTION_INITIALE = PREPARATION

_HEURE_LISIBLE = re.compile(r"^\d{1,2}\s*[hH:]\s*\d{2}$|^\d{3,4}$")


class ClasseurIllisible(ValueError):
    """Le fichier n'a pas la forme d'un rapport EOD.

    L'écran doit le dire plutôt qu'annoncer « 0 ligne importée », qui laisserait chercher du côté
    du contenu ce qui manque à la forme.
    """


@dataclass(frozen=True)
class LigneClasseur:
    """Une ligne du rapport, lue telle quelle. Rien n'y est encore rapproché du déroulé."""

    section: str
    libelle: str
    debut: str | None
    fin: str | None
    #: Les lignes « System Date » ne portent pas une heure mais une **date lue à l'écran**.
    valeur: date | None
    note: str | None
    sans_objet: bool


@dataclass(frozen=True)
class ClasseurEod:
    """Ce qu'un classeur dit d'une nuit."""

    #: La journée comptable close, lue sur le premier « System Date » — celui d'AVANT la bascule.
    #: C'est la seule date fiable du fichier : le nom de la feuille et celui du fichier se
    #: contredisent déjà sur les exemplaires reçus (« EOD 13 10 2025.xlsx » pour une nuit de 2026).
    journee: date | None
    lignes: list[LigneClasseur]


def cle_intitule(valeur: Any) -> str:
    """Clé de rapprochement d'un intitulé : sans accents, sans casse, sans ponctuation.

    `normaliser` ne touche pas à la ponctuation, et c'est précisément là que le classeur et le
    modèle divergent (« CSSJOBBR--- AC-DAHOFF » contre « CSSJOBBR --- AC-DAHOFF »). On réduit donc
    tout ce qui n'est ni lettre ni chiffre à une espace.
    """
    texte = unicodedata.normalize("NFKD", str(valeur or "").replace("\xa0", " "))
    sans_accent = "".join(c for c in texte if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^0-9a-zA-Z]+", " ", sans_accent).lower().split())


def _texte(valeur: Any) -> str | None:
    """Cellule texte nettoyée — espaces insécables et retours à la ligne compris."""
    if valeur in (None, ""):
        return None
    propre = " ".join(str(valeur).replace("\xa0", " ").split())
    return propre or None


def _heure(valeur: Any) -> str | None:
    """L'heure telle qu'écrite (« 20H29 »). La résoudre en horodatage revient au domaine."""
    texte = _texte(valeur)
    if texte is None or not _HEURE_LISIBLE.match(texte):
        return None
    return texte


def _date(valeur: Any) -> date | None:
    """La date lue à l'écran du core banking, pour les lignes « System Date »."""
    if isinstance(valeur, datetime):
        return valeur.date()
    if isinstance(valeur, date):
        return valeur
    return None


def _reperer_entetes(feuille: Any) -> tuple[int, dict[str, int]] | None:
    """La ligne d'en-tête et la colonne de chaque champ, ou ``None`` si la feuille n'en a pas."""
    for ligne in feuille.iter_rows(min_row=1, max_row=min(feuille.max_row, 20)):
        trouvees = {
            normaliser(cellule.value): cellule.column
            for cellule in ligne
            if normaliser(cellule.value) in _COLONNES
        }
        if len(trouvees) == len(_COLONNES):
            # L'intitulé vit dans la colonne qui précède « Start » : le classeur n'en nomme pas
            # l'en-tête (il y écrit « AFG-BANK-MALI »), on la déduit donc de la position.
            trouvees["libelle"] = trouvees["start"] - 1
            return ligne[0].row, trouvees
    return None


def analyser_classeur(contenu: bytes) -> ClasseurEod:
    """Lit un classeur EOD et rend ses lignes, normalisées mais non interprétées.

    Lève ``ClasseurIllisible`` si aucune feuille ne porte les trois en-têtes du rapport.
    """
    classeur = openpyxl.load_workbook(BytesIO(contenu), data_only=True)

    repere = None
    for feuille in classeur.worksheets:
        repere = _reperer_entetes(feuille)
        if repere is not None:
            break
    if repere is None or feuille is None:
        raise ClasseurIllisible(
            "Ce fichier n'a pas la forme d'un rapport EOD : les colonnes « Start », « End » et "
            "« Notes » sont introuvables."
        )
    entete, colonnes = repere

    section = _SECTION_INITIALE
    lignes: list[LigneClasseur] = []
    journee: date | None = None

    for numero in range(entete + 1, feuille.max_row + 1):
        libelle = _texte(feuille.cell(numero, colonnes["libelle"]).value)
        if libelle is None:
            continue
        brut_debut = feuille.cell(numero, colonnes["start"]).value
        brut_fin = feuille.cell(numero, colonnes["end"]).value
        brut_note = feuille.cell(numero, colonnes["notes"]).value

        # Une ligne de section ne porte qu'un intitulé : le classeur la fusionne sur toute la
        # largeur. C'est ce vide qui la distingue, pas son nom — un nouveau « PART 5 » se rangerait
        # tout seul, là où une liste figée l'aurait rejeté.
        if all(v in (None, "") for v in (brut_debut, brut_fin, brut_note)):
            section = _SECTIONS_TRADUITES.get(cle_intitule(libelle), libelle)
            continue

        sans_objet = any(
            normaliser(v) in _SANS_OBJET for v in (brut_debut, brut_fin, brut_note) if v is not None
        )
        valeur = _date(brut_debut)
        note = _texte(brut_note)
        if note is not None and (
            cle_intitule(note) in _VERDICTS_MUETS or normaliser(note) in _SANS_OBJET
        ):
            note = None

        lignes.append(
            LigneClasseur(
                section=section,
                libelle=libelle,
                debut=None if valeur is not None else _heure(brut_debut),
                fin=None if valeur is not None else _heure(brut_fin),
                valeur=valeur,
                note=note,
                sans_objet=sans_objet,
            )
        )
        # Le premier « System Date » est relevé AVANT la bascule : il porte la journée qu'on clôt.
        if journee is None and valeur is not None:
            journee = valeur

    if not lignes:
        raise ClasseurIllisible("Le rapport ne contient aucune ligne d'étape.")
    return ClasseurEod(journee=journee, lignes=lignes)
